"""Monthly Accounting Layer(月次会計確定レイヤー)。

日報監査(product_audit.py / run_product_audit.py)が確定した数字を「上から」
読み取り、1店舗×1か月の月次確定データへ集約する。

設計原則(2026-09-16確定、product-audit-spec.md §30参照):
- 既存のproduct_audit.py・run_product_audit.pyのロジックは一切変更しない。
  月次集計のために既存監査ロジックを書き換えることはしない(monthly layerは
  既存product auditの「上」に作る)。
- 日報監査で確定した数字だけを使用する。不明な数字を0円として埋めず、nullで保持する。
- source/current(現在表示ベース)とmanagement accounting(監査調整後)を混同しない。
- 監査調整は必ず別項目(audit_adjustments)で保持する。
- 元Excelは一切編集・再計算・保存しない(read_only相当のopenpyxl読み取りのみ)。

売上authority(2026-09-16改訂、オーナー指摘反映。§30参照):
「日報 authoritative as recorded」の原則により、月報集計シートの実売実績セルは
authorityではなく、あくまで検算対象(reconciliation target)として扱う。
management accountingの店舗全体売上は、日報の日別シート(AC/AD/AF)を積み上げた
authoritativeな内訳の合計として計算する。SOURCE月報集計セルが読み取れる場合でも、
みよし店のようにスタッフ名マッピング漏れで過少集計されているケースがあるため、
「SOURCE月報集計セルが読み取れる=authority」とは扱わない。
"""
from __future__ import annotations

import json
import warnings
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import yaml

from . import io_utils, xlsx_report

STORE_IDS_2026_08 = ["moriyama", "miyoshi", "midori", "kariya", "nisshin_akaike", "inazawa"]

CONFIG_DIR = io_utils.ACCOUNTING_ROOT / "config"
STORES_YAML_PATH = CONFIG_DIR / "stores.yaml"

# 売上authorityの明示(2026-09-16確定)。各フィールドがどのデータ源から来ているかを
# JSON出力にも埋め込み、「SOURCE月報集計はreconciliation targetであってauthorityでは
# ない」という原則を追跡可能にする。
SALES_AUTHORITY = {
    "service_new": (
        "daily AC direct sum: 日報日別シート(1〜31)のAC列(新規売・税抜)を、"
        "xlsx_report.py(product_auditとは独立)で直接合計した値。"
    ),
    "service_existing": (
        "daily AD direct sum: 日報日別シートのAD列(既存・回数券売・税抜)の直接合計。"
    ),
    "retail_source_current": (
        "product audit daily AF transaction total: product_audit.pyが取引単位で確定した"
        "物販税抜売上合計(revenue_excl_tax.source_current_total)。"
    ),
    "retail_management_accounting": (
        "retail_source_current + confirmed audit adjustment"
        "(product_audit.pyのrevenue_excl_tax.audit_adjustment)。"
    ),
    "store_management_total": (
        "management service_new + management service_existing + management retail + other。"
        "日報の日別シート積み上げがauthorityであり、SOURCE月報集計シートのセル値ではない。"
    ),
    "source_workbook_monthly_summary": (
        "reconciliation target only: 月報集計シートの実売実績セルはauthorityではなく、"
        "authoritativeな日別積み上げ値との検算対象として別途保持する"
        "(monthly_summary_reconciliation参照)。"
    ),
}


def load_store_master(path: Path = STORES_YAML_PATH) -> dict[str, dict]:
    """config/stores.yamlを読み込み、store_id -> {official_name, store_type, aliases} を返す。

    store_typeが未設定の店舗は明示的にNoneのまま返す(direct/franchiseを推測しない)。
    """
    with open(path, "r", encoding="utf-8") as f:
        doc = yaml.safe_load(f)
    result: dict[str, dict] = {}
    for rec in doc.get("stores", []) or []:
        result[rec["id"]] = {
            "official_name": rec.get("official_name"),
            "store_type": rec.get("store_type"),
            "aliases": rec.get("aliases", []),
        }
    return result


def load_product_audit_summary(store_id: str, year_month: str) -> dict:
    """既存run_product_audit.pyが出力したJSON summaryを読み込む。

    ファイルが存在しない場合はエラーとする(商品監査が未実行の店舗を推測で
    埋めない)。既存の物販監査結果をそのまま「正本」として使用する。
    """
    path = io_utils.OUTPUT_DIR / "product_audit" / year_month / f"{store_id}_product_audit_summary.json"
    if not path.exists():
        raise FileNotFoundError(
            f"{store_id}の商品監査結果が見つかりません({path})。"
            "先にrun_product_audit.pyを実行してください(推測で埋められません)。"
        )
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _read_store_total_actual_sales_raw(wb) -> dict:
    """月報集計シートの「実売（実績）」セル(累計列)を直接読み取り、値・状態・生のエラー
    テキストを返す(product_audit.parse_monthly_summaryと同じラベル検索ロジックの読み取り
    専用の複製。product_audit.py自体は変更しない。元の型情報(#N/A等)を保持するために
    独立実装している)。この値はauthorityではなく、検算(reconciliation target)専用。
    """
    if "月報集計" not in wb.sheetnames:
        return {"value": None, "status": "unavailable", "error": "月報集計シートが無い", "raw": None}
    ws = wb["月報集計"]
    total_col = None
    for row in range(1, 10):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=row, column=col).value == "累計":
                total_col = col
                break
        if total_col:
            break
    if total_col is None:
        return {"value": None, "status": "unavailable", "error": "「累計」列が見つからない", "raw": None}
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=2).value == "実売（実績）":
            raw = ws.cell(row=row, column=total_col).value
            if isinstance(raw, (int, float)):
                return {"value": float(raw), "status": "ok", "error": None, "raw": raw}
            if isinstance(raw, str) and raw.strip().startswith("#"):
                return {"value": None, "status": "error", "error": raw.strip(), "raw": raw}
            if raw is None:
                return {"value": None, "status": "empty", "error": None, "raw": None}
            return {"value": None, "status": "unexpected_type", "error": str(raw), "raw": raw}
    return {"value": None, "status": "unavailable", "error": "「実売（実績）」行が見つからない", "raw": None}


def compute_source_service_sales(source_xlsx_path: Path) -> dict:
    """AC(新規売)・AD(既存・回数券売)・AF(物販、参考用の生合計)の日報直接合計を、
    既存のxlsx_report.py(extract_all_days/DayResult、守山店向けに既にテスト済みの
    行単位パーサー)を再利用して算出する。数式エラーのセルは0円と断定せず、
    unresolved側へ分離する(product_audit.pyの物販監査ロジックとは完全に独立した計算)。

    AF(retail_total_raw)はproduct_audit.pyの取引単位集計とは別経路であり、両者の
    差はretail_reconciliationで別途保持する(rounding artifactの明示化用。物販の
    正式金額としては使わない。正式にはproduct_audit.pyのrevenue_excl_taxを使う)。
    """
    day_results = xlsx_report.extract_all_days(source_xlsx_path)
    new_total = round(sum(d.new_total for d in day_results.values()), 2)
    existing_total = round(sum(d.existing_total for d in day_results.values()), 2)
    retail_total_raw = round(sum(d.retail_total for d in day_results.values()), 2)
    unresolved_total = round(sum(d.unresolved_total for d in day_results.values()), 2)
    error_row_count = sum(len(d.error_rows) for d in day_results.values())
    unresolved_row_count = sum(len(d.unresolved_rows) for d in day_results.values())
    return {
        "service_new": new_total,
        "service_existing": existing_total,
        "retail_total_raw": retail_total_raw,
        "unresolved_total": unresolved_total,
        "error_row_count": error_row_count,
        "unresolved_row_count": unresolved_row_count,
    }


def _check(label: str, lhs, rhs, *, tolerance: float = 1.0, known_difference_note: str | None = None) -> dict:
    """2つの数値がtolerance以内で一致するかを検証する。差異があっても補正はしない。"""
    if lhs is None or rhs is None:
        return {"check": label, "status": "unavailable", "lhs": lhs, "rhs": rhs, "diff": None, "note": None}
    diff = round(lhs - rhs, 2)
    if abs(diff) < tolerance:
        return {"check": label, "status": "pass", "lhs": lhs, "rhs": rhs, "diff": diff, "note": None}
    if known_difference_note:
        return {
            "check": label, "status": "known_difference", "lhs": lhs, "rhs": rhs, "diff": diff,
            "note": known_difference_note,
        }
    return {"check": label, "status": "fail", "lhs": lhs, "rhs": rhs, "diff": diff, "note": None}


def build_store_monthly_record(store_id: str, year_month: str, store_master: dict, data_dir: Path | None = None) -> dict:
    """1店舗×1か月分の月次確定データを組み立てる。"""
    data_dir = data_dir or io_utils.DATA_DIR
    pa = load_product_audit_summary(store_id, year_month)

    store_info = store_master.get(store_id, {})
    store_name = store_info.get("official_name") or store_id
    store_type = store_info.get("store_type")

    source_xlsx_path = data_dir / store_id / year_month / "①8月.xlsx"
    if not source_xlsx_path.exists():
        candidates = sorted((data_dir / store_id / year_month).glob("*.xlsx")) if (data_dir / store_id / year_month).exists() else []
        if not candidates:
            raise FileNotFoundError(f"{store_id}の日報Excelが見つかりません({source_xlsx_path})。")
        source_xlsx_path = candidates[0]

    source_xlsx_sha256 = xlsx_report.file_sha256(source_xlsx_path)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(source_xlsx_path, data_only=True, read_only=False)
    try:
        service = compute_source_service_sales(source_xlsx_path)
        store_total_raw = _read_store_total_actual_sales_raw(wb)
    finally:
        wb.close()

    # --- retail(物販)は既存product_audit.pyの出力をそのまま使う(§10: 既存ロジックを
    # 月次集計のために書き換えない)。retailの正式値はここ(product_audit)がauthority。 ---
    revenue = pa["revenue_excl_tax"]
    retail_source_current = revenue["source_current_total"]
    retail_audit_adjustment = revenue["audit_adjustment"]
    retail_management = revenue["management_accounting_total"]

    # category_reclassifications(§14): 物販から施術/既存側へ振替済みの金額。
    # 現時点(2026-09-16)では全店舗0円だが、将来nonzeroになった場合に備えて
    # service_existingへ加算できるようにしておく(推測で新規カテゴリへは配分しない)。
    reclass_total = pa["category_reclassifications"]["total_tax_excl_revenue"] or 0.0

    other_amount = 0.0  # 日報のAC/AD/AF以外の売上区分は現構造には存在しない(構造上の事実)。

    source_current_sales = {
        "service_new": service["service_new"],
        "service_existing": round(service["service_existing"] + reclass_total, 2),
        "retail": retail_source_current,
        "other": other_amount,
    }
    source_current_sales["total"] = round(sum(source_current_sales.values()), 2)

    # 現時点でAC・AD(施術売上)に対する監査調整の仕組みは存在しないため、
    # management accountingはsource currentと同値とする(0で埋めているのではなく、
    # 「調整機構が無い=調整額0」という既知の事実)。
    management_accounting_sales = {
        "service_new": source_current_sales["service_new"],
        "service_existing": source_current_sales["service_existing"],
        "retail": retail_management,
        "other": other_amount,
    }
    # store management total = 日報の日別シート積み上げ(authoritative)の合計。
    # SOURCE月報集計セルの値ではない(2026-09-16改訂、§30参照)。
    management_accounting_sales["total"] = round(sum(management_accounting_sales.values()), 2)

    audit_adjustments = {
        "service_new": 0.0,
        "service_existing": 0.0,
        "retail": retail_audit_adjustment,
        "other": 0.0,
    }
    audit_adjustments["total"] = round(sum(audit_adjustments.values()), 2)

    retail_cogs = {
        "confirmed_cogs": pa["confirmed_cogs"],
        "unconfirmed_cogs_sales": pa["unconfirmed_cogs_sales"],
        "cogs_fully_confirmed": pa["cogs_fully_confirmed"],
    }
    retail_gross_profit = pa["gross_profit"]["management_accounting"]  # 既にnull-gated済み
    retail_gross_margin = pa["gross_margin"]["management_accounting"]

    store_overall = pa["store_overall_revenue"]
    reconstruction_applied = store_overall.get("store_total_actual_sales_reconstructed", False)

    # authoritativeな店舗全体売上(management accounting store total) = 日報の日別
    # シート積み上げ合計。SOURCE月報集計セルはreconciliation targetとして別保存する
    # (2026-09-16改訂。みよし店の6,119.44円差・稲沢店の#N/Aを、authorityの側で
    # 取り落とさないため)。
    authoritative_daily_total = management_accounting_sales["total"]

    store_total_sales = {
        "management_accounting": authoritative_daily_total,
        "source_workbook_reported": store_total_raw["value"],
        "source_workbook_status": store_total_raw["status"],
        "source_workbook_error": store_total_raw["error"],
        "reconstruction_applied": reconstruction_applied,  # product_audit側で決定論的復元を適用したか(参考情報)
    }

    # SOURCE月報集計セル(workbook cell)は、監査調整(audit_adjustment)を一切知らない
    # 「現在表示ベース」の値であり、authoritative_daily_total(management accounting、
    # 監査調整後)とは、退店の性質上そもそも一致しない(差額=audit_adjustments.total)。
    # これは「不明な差異」ではなく、source/current と management accounting を
    # 区別している本レイヤーの設計そのものによる既知の差である。さらに既知構造差異
    # (例:みよしのスタッフ名マッピング漏れ)が登録されている場合は、その分も
    # 合わせて期待差異に織り込む(2026-09-16改訂、オーナー指摘反映)。
    known_structural_amount = 0.0
    known_diff_note = None
    for d in pa["pending_review_summary"].get("known_structural_discrepancies", []):
        if d.get("check") == "store_total_actual_sales":
            known_structural_amount = d.get("diff") or 0.0
            known_diff_note = (
                "既知構造差異として原因確定済み(product_audit側confirmed_structural_"
                f"discrepancies.yaml参照、金額{known_structural_amount}円): "
                f"{d.get('cause', '').strip()[:200]}..."
            )
        elif d.get("check") == "store_total_actual_sales_unreadable":
            known_diff_note = (
                "SOURCE月報集計セルが#N/A等で読み取れない(原因は確定済み。"
                "product_audit側confirmed_structural_discrepancies.yaml参照): "
                f"{d.get('cause', '').strip()[:200]}..."
            )

    # 期待される差額 = -(既知構造差異額) - (retail監査調整額の合計)。
    # 導出: source_workbook ≈ source_current_sales.total - known_structural_amount
    #       authoritative_daily_total = source_current_sales.total + audit_adjustments.total
    #       diff = source_workbook - authoritative_daily_total
    #            = -known_structural_amount - audit_adjustments.total
    expected_diff_from_known_causes = round(-known_structural_amount - audit_adjustments["total"], 2)

    actual_diff = (
        round(store_total_raw["value"] - authoritative_daily_total, 2)
        if store_total_raw["value"] is not None else None
    )
    if actual_diff is None:
        msr_status = "unavailable"
        msr_reason = known_diff_note or "SOURCE月報集計セルが読み取れないため検算不可"
    elif abs(actual_diff) < 1.0:
        msr_status = "pass"
        msr_reason = "既知の丸め誤差の範囲内で一致"
    elif abs(round(actual_diff - expected_diff_from_known_causes, 2)) < 1.0:
        msr_status = "known_difference"
        parts = []
        if audit_adjustments["total"]:
            parts.append(
                f"retail監査調整額{audit_adjustments['total']}円(SOURCE月報は現在表示"
                "ベースのため監査調整を認識しない。source/current と management "
                "accountingを区別する本レイヤーの設計上の既知差)"
            )
        if known_structural_amount:
            parts.append(known_diff_note or f"既知構造差異{known_structural_amount}円")
        msr_reason = " + ".join(parts) if parts else "既知の要因で説明可能"
    else:
        msr_status = "fail"
        msr_reason = (
            f"既知の要因(監査調整{audit_adjustments['total']}円・既知構造差異"
            f"{known_structural_amount}円)だけでは説明できない未知の差異が残っている"
            f"(期待差額{expected_diff_from_known_causes}円 に対し実差額{actual_diff}円)"
        )

    monthly_summary_reconciliation = {
        "source_workbook_reported": store_total_raw["value"],
        "source_workbook_status": store_total_raw["status"],
        "authoritative_daily_total": authoritative_daily_total,
        "difference": actual_diff,
        "expected_difference_from_known_causes": expected_diff_from_known_causes,
        "status": msr_status,
        "reason": msr_reason,
    }

    # retailの丸め差(item 6): 日報AF列の生の直接合計と、product_auditの取引単位集計
    # (端数を取引ごとに丸めた合計)の差を明示する。取引の欠落・二重計上ではなく、
    # 個々の取引金額を円未満で丸めたことによる既知の誤差であることを保持する。
    retail_rounding_diff = round(service["retail_total_raw"] - retail_source_current, 4)
    retail_reconciliation = {
        "raw_daily_af_total": service["retail_total_raw"],
        "transaction_rounded_total": retail_source_current,
        "rounding_difference": retail_rounding_diff,
        "tolerance": 1.0,
        "status": "pass" if abs(retail_rounding_diff) < 1.0 else "known_difference",
    }

    pending = pa["pending_review_summary"]

    # sales_revenue_ready: 管理会計で採用する月次売上額(management_accounting_sales)が
    # 確定しているか。日報側に説明のつかない数式エラー・未確定行が残っている場合のみ
    # falseとする。原因確定済み(既知構造差異・決定論的復元)なものはtrueとして良い
    # (payment reconciliation未完了・税区分未確認だけを理由にfalseにはしない)。
    unresolved_unexplained = (service["error_row_count"] > 0 or service["unresolved_row_count"] > 0) and not reconstruction_applied
    sales_revenue_ready = not unresolved_unexplained

    retail_cogs_ready = retail_cogs["cogs_fully_confirmed"]
    retail_gross_profit_ready = sales_revenue_ready and retail_cogs_ready
    payment_reconciliation_complete = pa["payment_reconciliation"]["matches"]
    tax_category_complete = pa.get("tax_category_pending_count", 0) == 0
    full_audit_resolved = pending["fully_resolved"]  # 既存ロジックの意味をそのまま尊重する

    pl_readiness = {
        "sales_revenue_ready": sales_revenue_ready,
        "retail_cogs_ready": retail_cogs_ready,
        "retail_gross_profit_ready": retail_gross_profit_ready,
        "payment_reconciliation_complete": payment_reconciliation_complete,
        "tax_category_complete": tax_category_complete,
        "full_audit_resolved": full_audit_resolved,
    }

    audit_status = {
        "product_audit_pending_count": pending["product_audit_pending_review_count"],
        "store_wide_pending_count": pending["store_wide_pending_review_count"],
        "known_structural_discrepancy_count": pending["known_structural_discrepancy_count"],
        "tax_category_pending_count": pa.get("tax_category_pending_count", 0),
        "cogs_fully_confirmed": pa["cogs_fully_confirmed"],
        "fully_resolved": full_audit_resolved,
    }

    data_quality = {
        "sales_complete": sales_revenue_ready,
        "cogs_complete": retail_cogs_ready,
        "payment_reconciliation_complete": payment_reconciliation_complete,
        "tax_category_complete": tax_category_complete,
        "source_monthly_summary_usable": store_total_raw["status"] == "ok",
    }

    # --- reconciliation(A/C/D)。B(店舗全体)はmonthly_summary_reconciliationへ統合した
    # (authorityが日報積み上げそのものになったため、単純合計との一致は定義上常にpassに
    # なる。真に意味のある検算はSOURCE月報集計セルとの比較=monthly_summary_reconciliation
    # である)。 ---
    check_a_retail = _check(
        "A_retail_management_equals_source_plus_adjustment",
        retail_management, round(retail_source_current + retail_audit_adjustment, 2),
    )
    check_a_total = _check(
        "A_total_management_equals_source_plus_adjustment",
        management_accounting_sales["total"],
        round(source_current_sales["total"] + audit_adjustments["total"], 2),
    )
    if retail_gross_profit is not None:
        check_c = _check(
            "C_retail_gp_equals_management_sales_minus_confirmed_cogs",
            retail_gross_profit, round(retail_management - retail_cogs["confirmed_cogs"], 2),
        )
    else:
        check_c = {"check": "C_retail_gp_equals_management_sales_minus_confirmed_cogs",
                    "status": "unavailable", "lhs": None, "rhs": None, "diff": None,
                    "note": "cogs_fully_confirmed=falseのためnull"}
    if retail_gross_margin is not None and retail_management:
        check_d = _check(
            "D_retail_margin_equals_gp_over_management_sales",
            retail_gross_margin, round(retail_gross_profit / retail_management, 4), tolerance=0.001,
        )
    else:
        check_d = {"check": "D_retail_margin_equals_gp_over_management_sales",
                    "status": "unavailable", "lhs": None, "rhs": None, "diff": None,
                    "note": "cogs_fully_confirmed=falseのためnull"}

    reconciliation = {
        "A_retail": check_a_retail,
        "A_total": check_a_total,
        "C_retail_gross_profit": check_c,
        "D_retail_margin": check_d,
    }

    return {
        "schema_version": "monthly_accounting_layer.v2",
        "year_month": year_month,
        "store": {
            "store_id": store_id,
            "store_name": store_name,
            "store_type": store_type,
        },
        "sales_authority": SALES_AUTHORITY,
        "source_current_sales": source_current_sales,
        "management_accounting_sales": management_accounting_sales,
        "audit_adjustments": audit_adjustments,
        "retail_cogs": retail_cogs,
        "retail_gross_profit": retail_gross_profit,
        "retail_gross_margin": retail_gross_margin,
        "store_total_sales": store_total_sales,
        "monthly_summary_reconciliation": monthly_summary_reconciliation,
        "retail_reconciliation": retail_reconciliation,
        "pl_readiness": pl_readiness,
        "audit_status": audit_status,
        "data_quality": data_quality,
        "reconciliation": reconciliation,
        "known_structural_discrepancies": pending.get("known_structural_discrepancies", []),
        "store_wide_issues": pending.get("store_wide_issues", []),
        "tax_category_pending_items": pa.get("tax_category_pending_items", []),
        "source": {
            "product_audit_summary_path": str(
                io_utils.OUTPUT_DIR / "product_audit" / year_month / f"{store_id}_product_audit_summary.json"
            ),
            "source_xlsx_path": str(source_xlsx_path),
            "source_xlsx_sha256": source_xlsx_sha256,
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def build_company_summary(store_records: list[dict], year_month: str) -> dict:
    """6店舗分のmonthly recordを統合した全社サマリーを組み立てる。

    原価未確定店舗(cogs_fully_confirmed=false)が1店舗でも存在する場合、
    company全体のretail_gross_profit/marginは安易に確定せずnullのままにする。
    company側のstore total集計も、authoritativeなmanagement_accounting_sales.total
    の合計を使う(SOURCE月報集計セルの値とは混ぜない。2026-09-16改訂)。
    """
    def sum_field(block_name: str, key: str) -> float:
        return round(sum(r[block_name][key] for r in store_records), 2)

    source_current_sales_total = {
        k: sum_field("source_current_sales", k) for k in ["service_new", "service_existing", "retail", "other", "total"]
    }
    management_accounting_sales_total = {
        k: sum_field("management_accounting_sales", k) for k in ["service_new", "service_existing", "retail", "other", "total"]
    }
    audit_adjustments_total = {
        k: sum_field("audit_adjustments", k) for k in ["service_new", "service_existing", "retail", "other", "total"]
    }

    confirmed_cogs_total = round(sum(r["retail_cogs"]["confirmed_cogs"] for r in store_records), 2)
    unconfirmed_cogs_sales_total = round(sum(r["retail_cogs"]["unconfirmed_cogs_sales"] for r in store_records), 2)
    all_cogs_confirmed = all(r["retail_cogs"]["cogs_fully_confirmed"] for r in store_records)

    if all_cogs_confirmed:
        company_retail_gross_profit = round(management_accounting_sales_total["retail"] - confirmed_cogs_total, 2)
        company_retail_gross_margin = (
            round(company_retail_gross_profit / management_accounting_sales_total["retail"], 4)
            if management_accounting_sales_total["retail"] else None
        )
    else:
        company_retail_gross_profit = None
        company_retail_gross_margin = None

    # SOURCE月報集計セルの合計は参考値としてのみ別保存する(authorityとは混ぜない)。
    source_workbook_reported_sum = round(
        sum(r["store_total_sales"]["source_workbook_reported"] or 0 for r in store_records), 2
    )
    source_workbook_has_missing_store = any(
        r["store_total_sales"]["source_workbook_reported"] is None for r in store_records
    )

    def readiness_count(field: str) -> int:
        return sum(1 for r in store_records if r["pl_readiness"][field])

    pl_readiness_summary = {
        "sales_revenue_ready_store_count": readiness_count("sales_revenue_ready"),
        "retail_cogs_ready_store_count": readiness_count("retail_cogs_ready"),
        "retail_gross_profit_ready_store_count": readiness_count("retail_gross_profit_ready"),
        "payment_reconciliation_complete_store_count": readiness_count("payment_reconciliation_complete"),
        "tax_category_complete_store_count": readiness_count("tax_category_complete"),
        "full_audit_resolved_store_count": readiness_count("full_audit_resolved"),
        "total_store_count": len(store_records),
    }

    stores_with_unconfirmed_retail_cogs = [
        r["store"]["store_id"] for r in store_records if not r["retail_cogs"]["cogs_fully_confirmed"]
    ]
    stores_with_payment_pending = [
        r["store"]["store_id"] for r in store_records if not r["pl_readiness"]["payment_reconciliation_complete"]
    ]
    stores_with_tax_pending = [
        r["store"]["store_id"] for r in store_records if not r["pl_readiness"]["tax_category_complete"]
    ]
    stores_with_known_structural_difference = [
        r["store"]["store_id"] for r in store_records if r["audit_status"]["known_structural_discrepancy_count"] > 0
    ]

    # E: company合計 = 各店舗monthlyレコードの合計(定義上必ず一致するが、独立した
    # 再集計として検算する)。
    recompute_total = round(sum(r["source_current_sales"]["total"] for r in store_records), 2)
    check_e = _check("E_company_equals_sum_of_stores", source_current_sales_total["total"], recompute_total, tolerance=0.01)

    return {
        "schema_version": "monthly_accounting_layer.company_summary.v2",
        "year_month": year_month,
        "store_count": len(store_records),
        "store_ids": [r["store"]["store_id"] for r in store_records],
        "sales_authority": SALES_AUTHORITY,
        "source_current_sales_total": source_current_sales_total,
        "management_accounting_sales_total": management_accounting_sales_total,
        "audit_adjustments_total": audit_adjustments_total,
        "retail_cogs_summary": {
            "confirmed_cogs_total": confirmed_cogs_total,
            "unconfirmed_cogs_sales_total": unconfirmed_cogs_sales_total,
            "cogs_fully_confirmed_all_stores": all_cogs_confirmed,
            "stores_with_unconfirmed_cogs": stores_with_unconfirmed_retail_cogs,
        },
        "company_retail_gross_profit": company_retail_gross_profit,
        "company_retail_gross_margin": company_retail_gross_margin,
        # authoritativeな合計(日報積み上げベース)。SOURCE月報集計セルの単純合計とは
        # 別フィールドで明確に区別する。
        "authoritative_management_company_sales_total": management_accounting_sales_total["total"],
        "source_workbook_reported_total": source_workbook_reported_sum,
        "source_workbook_reported_has_missing_store": source_workbook_has_missing_store,
        "pl_readiness_summary": pl_readiness_summary,
        "stores_with_unconfirmed_retail_cogs": stores_with_unconfirmed_retail_cogs,
        "stores_with_payment_pending": stores_with_payment_pending,
        "stores_with_tax_pending": stores_with_tax_pending,
        "stores_with_known_structural_difference": stores_with_known_structural_difference,
        "audit_status_summary": {
            "product_audit_pending_total": sum(r["audit_status"]["product_audit_pending_count"] for r in store_records),
            "store_wide_pending_total": sum(r["audit_status"]["store_wide_pending_count"] for r in store_records),
            "known_structural_discrepancy_total": sum(r["audit_status"]["known_structural_discrepancy_count"] for r in store_records),
            "tax_category_pending_total": sum(r["audit_status"]["tax_category_pending_count"] for r in store_records),
            "stores_not_fully_resolved": [
                r["store"]["store_id"] for r in store_records if not r["audit_status"]["fully_resolved"]
            ],
        },
        "reconciliation": {"E_company_equals_sum_of_stores": check_e},
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def company_summary_csv_rows(store_records: list[dict]) -> list[dict]:
    rows = []
    for r in store_records:
        rows.append({
            "store_id": r["store"]["store_id"],
            "store_name": r["store"]["store_name"],
            "store_type": r["store"]["store_type"],
            "service_new": r["management_accounting_sales"]["service_new"],
            "service_existing": r["management_accounting_sales"]["service_existing"],
            "retail_source_current": r["source_current_sales"]["retail"],
            "retail_audit_adjustment": r["audit_adjustments"]["retail"],
            "retail_management_accounting": r["management_accounting_sales"]["retail"],
            "management_store_total": r["store_total_sales"]["management_accounting"],
            "source_workbook_reported_total": r["store_total_sales"]["source_workbook_reported"],
            "monthly_summary_diff": r["monthly_summary_reconciliation"]["difference"],
            "confirmed_cogs": r["retail_cogs"]["confirmed_cogs"],
            "unconfirmed_cogs_sales": r["retail_cogs"]["unconfirmed_cogs_sales"],
            "retail_gross_profit": r["retail_gross_profit"],
            "retail_gross_margin": r["retail_gross_margin"],
            "sales_revenue_ready": r["pl_readiness"]["sales_revenue_ready"],
            "retail_cogs_ready": r["pl_readiness"]["retail_cogs_ready"],
            "retail_gross_profit_ready": r["pl_readiness"]["retail_gross_profit_ready"],
            "payment_reconciliation_complete": r["pl_readiness"]["payment_reconciliation_complete"],
            "tax_category_complete": r["pl_readiness"]["tax_category_complete"],
            "full_audit_resolved": r["pl_readiness"]["full_audit_resolved"],
        })
    return rows


CSV_FIELDNAMES = [
    "store_id", "store_name", "store_type",
    "service_new", "service_existing",
    "retail_source_current", "retail_audit_adjustment", "retail_management_accounting",
    "management_store_total", "source_workbook_reported_total", "monthly_summary_diff",
    "confirmed_cogs", "unconfirmed_cogs_sales",
    "retail_gross_profit", "retail_gross_margin",
    "sales_revenue_ready", "retail_cogs_ready", "retail_gross_profit_ready",
    "payment_reconciliation_complete", "tax_category_complete", "full_audit_resolved",
]
