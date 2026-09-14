#!/usr/bin/env python3
"""物販取引の価格・原価監査スクリプト(読み取り専用)。

やること:
  1. Google Driveミラーとして取得済みのローカルExcel(日報)を読み取る(書き込みはしない)。
  2. 日別シート(1〜31)から、物販(商品購入)取引をすべて抽出する。
  3. accounting/config/product_price_history.yaml・staff_price_rules.yaml を使い、
     「取引日時点で有効だった価格・原価」で監査する(現在価格での再計算はしない)。
  4. 一般顧客／スタッフ／不明を判定し、それぞれの売上・原価・粗利益・粗利率を集計する。
  5. 異常(価格不一致・原価未登録・価格履歴不足・商品不明・要現場確認等)を一覧化する。
  6. 結果を accounting/output/ へ、ログを accounting/logs/ へ出力する。

やらないこと:
  - Googleスプレッドシートへの書き込み。
  - Dropbox・Google Driveミラー・元Excelファイルの変更。
  - 異常を「不正」「改ざん」と自動判定すること(あくまで不一致の抽出)。
"""
from __future__ import annotations

import argparse
import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl

from lib import io_utils, product_audit, xlsx_report

DATA_START_ROW = xlsx_report.DATA_START_ROW
DATA_END_ROW = xlsx_report.DATA_END_ROW
OVERFLOW_START_ROW = xlsx_report.OVERFLOW_START_ROW
OVERFLOW_END_ROW = xlsx_report.OVERFLOW_END_ROW


def parse_args():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--store-id", default="moriyama")
    p.add_argument("--year-month", default="2026-08")
    p.add_argument(
        "--source-xlsx",
        default=str(io_utils.DATA_DIR / "moriyama" / "2026-08" / "①8月.xlsx"),
    )
    p.add_argument(
        "--price-history",
        default=str(Path(__file__).resolve().parent.parent / "config" / "product_price_history.yaml"),
    )
    p.add_argument(
        "--staff-price-rules",
        default=str(Path(__file__).resolve().parent.parent / "config" / "staff_price_rules.yaml"),
    )
    p.add_argument(
        "--staff-aliases",
        default=str(Path(__file__).resolve().parent.parent / "config" / "staff_aliases.yaml"),
    )
    p.add_argument(
        "--purchaser-blocks",
        default=str(Path(__file__).resolve().parent.parent / "config" / "confirmed_purchaser_blocks.yaml"),
    )
    p.add_argument(
        "--category-reclassifications",
        default=str(Path(__file__).resolve().parent.parent / "config" / "confirmed_category_reclassifications.yaml"),
    )
    p.add_argument(
        "--status-overrides",
        default=str(Path(__file__).resolve().parent.parent / "config" / "confirmed_status_overrides.yaml"),
    )
    p.add_argument(
        "--quantity-corrections",
        default=str(Path(__file__).resolve().parent.parent / "config" / "confirmed_quantity_corrections.yaml"),
    )
    p.add_argument(
        "--exception-transactions",
        default=str(Path(__file__).resolve().parent.parent / "config" / "confirmed_exception_transactions.yaml"),
    )
    return p.parse_args()


def _num(v):
    if v is None:
        return 0.0
    if isinstance(v, str):
        return None if v.strip().startswith("#") else 0.0
    return float(v)


def get_staff_names(wb) -> list[str]:
    ws = wb["基本情報"]
    names = []
    for row in range(4, 10):
        v = ws.cell(row=row, column=5).value
        if v and str(v).strip() not in ("その他",):
            names.append(str(v).strip())
    return names


def extract_retail_rows(ws, sheet: str):
    """物販(商品購入)取引の行を抽出する。書き込みは行わない。

    商品名(J列)が空欄・「無」であっても、物販売上(P列)が計上されている行は
    実在する取引として抽出する(商品名の有無を集計対象かどうかの条件にしない。
    ルール1〜3・11と同じ考え方)。この場合、商品名は「無」のまま監査に回し、
    価格履歴と突き合わせられないため H(商品不明)として検出させる。

    I列の生値も取得する。I列は単純な「備考」欄ではなく、「既存 単発・回数券」
    (施術チケットの種別)を記録する欄だが、特殊運用として「売上シェア」という
    文字列が入力される場合がある(2026-09-15オーナー確認、product-audit-spec.md §13)。
    守山8月の実データで、売上シェア分担入力の「無」側行にI列「売上シェア」の明示記載が
    あることを確認しており(product-audit-spec.md §11)、店舗・月をまたいで同様の記載が
    あればdetect_and_apply_shared_salesがそれを最優先の判定根拠として使う。
    記載が無い店舗・月では、この列は単にNoneのまま扱われ、判定には影響しない。"""
    rows = []
    for row in list(range(DATA_START_ROW, DATA_END_ROW + 1)) + list(range(OVERFLOW_START_ROW, OVERFLOW_END_ROW + 1)):
        j = ws[f"J{row}"].value
        k = _num(ws[f"K{row}"].value)
        p = _num(ws[f"P{row}"].value)
        if j in (None, "") and not k and not p:
            continue
        if not k and not p:
            continue
        t = _num(ws[f"T{row}"].value) or 0.0
        w = _num(ws[f"W{row}"].value) or 0.0
        af_raw = ws[f"AF{row}"].value
        # AF(税抜売上)自体が数式エラー(#N/A等)の場合、0円と断定せず「金額不明」として
        # 分離する(売上を落とさない・0円で処理しないという最重要ルールに基づく)。
        af_is_error = isinstance(af_raw, str) and af_raw.strip().startswith("#")
        af = 0.0 if af_is_error else (_num(af_raw) or 0.0)
        b = ws[f"B{row}"].value
        c = ws[f"C{row}"].value
        d = ws[f"D{row}"].value
        i_val = ws[f"I{row}"].value
        note_raw = str(i_val).strip() if i_val not in (None, "") else None
        rows.append({
            "row": row, "product": str(j).strip(), "quantity": k,
            "gross_incl_tax": p, "discount_incl_tax": t, "net_incl_tax": w,
            "tax_excl_revenue": af, "revenue_error": af_is_error,
            "customer_name": b, "staff_col": c, "category_label": d, "note_raw": note_raw,
            "purchaser_name_note": "", "quantity_correction_note": "",
        })
    return rows


def main():
    args = parse_args()
    store_id, year_month = args.store_id, args.year_month
    source_path = Path(args.source_xlsx)

    logger, log_path = io_utils.setup_logger("product_audit", store_id, year_month)
    logger.info("=== 物販監査スクリプト開始 ===")
    logger.info(f"対象: store_id={store_id}, year_month={year_month}")
    logger.info(f"入力ファイル(読み取り専用): {source_path}")

    if not source_path.exists():
        logger.error(f"入力ファイルが見つかりません: {source_path}")
        sys.exit(1)

    source_hash = xlsx_report.file_sha256(source_path)
    logger.info(f"入力ファイルのSHA-256: {source_hash}")

    price_history_path = Path(args.price_history)
    staff_rules_path = Path(args.staff_price_rules)
    staff_aliases_path = Path(args.staff_aliases)
    purchaser_blocks_path = Path(args.purchaser_blocks)
    category_reclass_path = Path(args.category_reclassifications)
    status_overrides_path = Path(args.status_overrides)
    quantity_corrections_path = Path(args.quantity_corrections)
    exception_transactions_path = Path(args.exception_transactions)
    price_history = product_audit.load_price_history(price_history_path)
    staff_price_rules = product_audit.load_staff_price_rules(staff_rules_path)
    staff_aliases = product_audit.load_staff_aliases(staff_aliases_path) if staff_aliases_path.exists() else {}
    purchaser_blocks = (
        product_audit.load_confirmed_purchaser_blocks(purchaser_blocks_path)
        if purchaser_blocks_path.exists() else []
    )
    category_reclass_map = (
        product_audit.load_confirmed_category_reclassifications(category_reclass_path)
        if category_reclass_path.exists() else {}
    )
    status_overrides = (
        product_audit.load_confirmed_status_overrides(status_overrides_path)
        if status_overrides_path.exists() else {}
    )
    quantity_corrections = (
        product_audit.load_confirmed_quantity_corrections(quantity_corrections_path)
        if quantity_corrections_path.exists() else {}
    )
    exception_transactions = (
        product_audit.load_confirmed_exception_transactions(exception_transactions_path)
        if exception_transactions_path.exists() else {}
    )
    logger.info(f"価格履歴マスター読み込み: {len(price_history)}商品 ({price_history_path})")
    logger.info(f"スタッフ価格履歴マスター読み込み: {len(staff_price_rules)}商品 ({staff_rules_path})")
    logger.info(f"スタッフ・社内購入者の別名マスター読み込み: {len(staff_aliases)}件 ({staff_aliases_path})")
    logger.info(f"確定済み連続購入ブロック読み込み: {len(purchaser_blocks)}件 ({purchaser_blocks_path})")
    logger.info(f"確定済み区分振替マスター読み込み: {len(category_reclass_map)}件 ({category_reclass_path})")
    logger.info(f"確定済みステータス上書きマスター読み込み: {len(status_overrides)}件 ({status_overrides_path})")
    logger.info(f"確定済み数量修正マスター読み込み: {len(quantity_corrections)}件 ({quantity_corrections_path})")
    logger.info(f"確定済み社内例外取引マスター読み込み: {len(exception_transactions)}件 ({exception_transactions_path})")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(source_path, data_only=True, read_only=False)

    staff_names = get_staff_names(wb)
    logger.info(f"担当スタッフ名一覧(購入者区分の判定に使用): {staff_names}")

    all_transactions: list[product_audit.ProductTransaction] = []
    all_reclassified_rows: list[dict] = []
    day_reconciliation_rows = []
    for day in xlsx_report.DAY_SHEETS:
        if day not in wb.sheetnames:
            continue
        ws = wb[day]
        date_key = f"{year_month}-{int(day):02d}"
        day_retail_rows = extract_retail_rows(ws, day)

        # 確定済み「同一購入者・連続購入ブロック」(§12): 顧客名が空欄の行に、
        # 起点行の購入者名を適用する。オーナー確認済みの範囲だけに限定する。
        product_audit.apply_confirmed_purchaser_blocks(
            day_retail_rows, store_id=store_id, year_month=year_month,
            date_key=date_key, blocks=purchaser_blocks,
        )

        # 確定済み数量修正(§19): 現場確認により実際の数量と異なると確定した行の
        # 数量を補正する(元Excelは変更しない)。
        product_audit.apply_confirmed_quantity_corrections(
            day_retail_rows, store_id=store_id, year_month=year_month,
            date_key=date_key, corrections=quantity_corrections,
        )

        # 確定済み区分振替(§14): 実際には物販ではなく施術・回数券(既存)側の売上と
        # 確定した行を、物販監査の対象から完全に除外する。
        kept_rows = []
        day_reclassified = []
        for r in day_retail_rows:
            rec = category_reclass_map.get((store_id, year_month, date_key, r["row"]))
            if rec is not None:
                day_reclassified.append({**r, **rec})
                logger.info(
                    f"[{day}日] 行{r['row']}: 物販監査から除外し、施術/既存側へ振替"
                    f"(税抜{rec['tax_excl_revenue']}円、{rec['reclassified_as']})"
                )
                continue
            kept_rows.append(r)
        day_retail_rows = kept_rows
        all_reclassified_rows.extend(day_reclassified)

        day_txs = []
        for r in day_retail_rows:
            tx = product_audit.audit_transaction(
                date=date_key, row=r["row"], customer_name=r["customer_name"],
                staff_col=r["staff_col"], category_label=r["category_label"],
                product_name=r["product"], quantity=r["quantity"],
                gross_incl_tax=r["gross_incl_tax"], discount_incl_tax=r["discount_incl_tax"],
                net_incl_tax=r["net_incl_tax"], tax_excl_revenue=r["tax_excl_revenue"],
                price_history=price_history, staff_price_rules=staff_price_rules,
                staff_names=staff_names, staff_aliases=staff_aliases, revenue_error=r["revenue_error"],
                note_raw=r["note_raw"], purchaser_name_note=r["purchaser_name_note"],
                quantity_correction_note=r["quantity_correction_note"],
            )
            day_txs.append(tx)
        all_transactions.extend(day_txs)

        # 調査1: 日別シートAF54(物販税別合計)と、監査エンジンがその日に認識した
        # 物販税抜売上の合計を突き合わせる。1円以上ずれた場合は行番号まで特定する。
        # AF54はD列(新規/既存/物販)を区別しない単純合計(=SUMIF(AF4:AF53,">0"))のため、
        # 区分振替(§14)で物販監査から除外した行の分だけ、監査エンジン側の合計より
        # 大きくなる。この既知の差額(day_reclass_total)を戻して比較する。
        af54_raw = ws["AF54"].value
        af54 = None if (isinstance(af54_raw, str) and af54_raw.startswith("#")) else (af54_raw or 0)
        engine_sum = round(sum(t.tax_excl_revenue for t in day_txs), 2)
        day_reclass_total = round(sum(r["tax_excl_revenue"] for r in day_reclassified), 2)
        diff = None if af54 is None else round(engine_sum + day_reclass_total - af54, 2)
        day_reconciliation_rows.append({
            "day": day, "af54_report": af54, "engine_sum": engine_sum,
            "reclassified_total": day_reclass_total, "diff": diff, "transaction_count": len(day_txs),
        })
        if diff is not None and abs(diff) >= 1.0:
            logger.error(
                f"[{day}日] 物販売上の不一致を検出: AF54={af54} / 監査エンジン={engine_sum} "
                f"/ 施術・既存へ振替={day_reclass_total} (差={diff})"
            )
            for r in day_retail_rows:
                logger.error(
                    f"  行{r['row']}: 商品={r['product']!r} 顧客={r['customer_name']!r} "
                    f"AF(税抜)={r['tax_excl_revenue']} P(税込)={r['gross_incl_tax']}"
                )

    logger.info(f"物販取引を{len(all_transactions)}件抽出しました。")
    if all_reclassified_rows:
        total_reclassified = round(sum(r["tax_excl_revenue"] for r in all_reclassified_rows), 2)
        logger.info(
            f"区分振替(物販から施術/既存へ): {len(all_reclassified_rows)}件・税抜合計{total_reclassified}円"
        )

    # 確定済みステータス上書き(§15・§19): 通常はclassification自体は変更せず、
    # severityのみ「分類保留」に変更する(事実は確定済みだが、社割ルール等のマスター登録が
    # 未了のもの)。現場確認により価格異常自体が解消したと確定できたがファイル内に対応する
    # 行が見当たらないケースに限り、classification自体もあわせて上書きする(§19)。
    overrides_applied = product_audit.apply_confirmed_status_overrides(
        all_transactions, store_id=store_id, year_month=year_month, overrides=status_overrides,
    )
    if overrides_applied:
        logger.info(f"確定済みステータス上書きを適用: {overrides_applied}件")

    # 確定済み社内例外取引(§18): 通常非販売の備品等を原価のまま社内購入した例外取引を、
    # 商品マスター・原価マスター不在を異常扱いせず、正常な例外取引として確定する。
    exceptions_applied = product_audit.apply_confirmed_exception_transactions(
        all_transactions, store_id=store_id, year_month=year_month, exceptions=exception_transactions,
    )
    if exceptions_applied:
        logger.info(f"確定済み社内例外取引を適用: {exceptions_applied}件(classification→N)")

    # 売上シェア(1商品を複数スタッフで分担入力)の検出。物販売上総額・元Excelは変更せず、
    # 該当取引のclassification・severity等のみ調整する(2026-09-15確定。
    # product-audit-spec.md §11参照)。
    share_groups = product_audit.detect_and_apply_shared_sales(all_transactions)
    confirmed_share_groups = [g for g in share_groups if g["transaction_type"] == "shared_sale"]
    candidate_share_groups = [g for g in share_groups if g["transaction_type"] == "shared_sale_candidate"]
    confirmed_2way_groups = [g for g in confirmed_share_groups if g["share_count"] == 2]
    confirmed_3way_groups = [g for g in confirmed_share_groups if g["share_count"] == 3]
    logger.info(
        f"売上シェア検出: 確定{len(confirmed_share_groups)}グループ"
        f"(2名{len(confirmed_2way_groups)}件・3名{len(confirmed_3way_groups)}件、"
        f"計{sum(g['share_count'] for g in confirmed_share_groups)}行) / "
        f"候補(未確定){len(candidate_share_groups)}グループ"
        f"({sum(g['share_count'] for g in candidate_share_groups)}行)"
    )
    for g in share_groups:
        logger.info(
            f"  [{g['transaction_type']}] {g['date']} 行{g['linked_rows']} "
            f"{g['linked_product']} 合計{g['share_total']}円 ({g['share_count']}名)"
        )

    reconciliation_out_dir = io_utils.OUTPUT_DIR / "product_audit" / year_month
    shared_sales_path = reconciliation_out_dir / f"{store_id}_shared_sales.csv"
    shared_sales_csv_rows = [
        {**g, "linked_rows": "・".join(str(r) for r in g["linked_rows"])} for g in share_groups
    ]
    io_utils.write_csv(
        shared_sales_path, shared_sales_csv_rows,
        fieldnames=["share_group_id", "transaction_type", "date", "linked_product",
                    "share_count", "share_total", "linked_rows"],
    )
    logger.info(f"売上シェア検出結果を出力: {shared_sales_path}")
    day_reconciliation_path = reconciliation_out_dir / f"{store_id}_af54_vs_engine.csv"
    io_utils.write_csv(
        day_reconciliation_path, day_reconciliation_rows,
        fieldnames=["day", "af54_report", "engine_sum", "reclassified_total", "diff", "transaction_count"],
    )
    logger.info(f"日別AF54突合結果を出力: {day_reconciliation_path}")

    category_reclass_path_out = reconciliation_out_dir / f"{store_id}_category_reclassifications.csv"
    category_reclass_csv_rows = [
        {
            "date": r["date"], "row": r["row"], "customer_name": r.get("customer_name") or "(空欄)",
            "staff_col": r.get("staff_col") or "", "product": r.get("product") or "",
            "tax_excl_revenue": r["tax_excl_revenue"], "reclassified_as": r["reclassified_as"],
            "reason": r["reason"].strip(),
        }
        for r in all_reclassified_rows
    ]
    io_utils.write_csv(
        category_reclass_path_out, category_reclass_csv_rows,
        fieldnames=["date", "row", "customer_name", "staff_col", "product",
                    "tax_excl_revenue", "reclassified_as", "reason"],
    )
    logger.info(f"区分振替(物販除外)一覧を出力: {category_reclass_path_out}({len(category_reclass_csv_rows)}件)")

    total_af54 = sum(r["af54_report"] or 0 for r in day_reconciliation_rows if r["af54_report"] is not None)
    total_engine = sum(r["engine_sum"] for r in day_reconciliation_rows)
    total_reclassified_all = round(sum(r["reclassified_total"] for r in day_reconciliation_rows), 2)
    total_diff = round(total_engine + total_reclassified_all - total_af54, 2)
    if abs(total_diff) >= 1.0:
        logger.error(
            f"月合計で不一致: SUM(AF54)={total_af54} / 監査エンジン合計={total_engine} "
            f"/ 施術・既存へ振替={total_reclassified_all} (差={total_diff})。この監査は完成扱いにできません。"
        )
    else:
        logger.info(
            f"月合計はSUM(AF54)={total_af54}と監査エンジン合計+区分振替額"
            f"({total_engine}+{total_reclassified_all}={round(total_engine + total_reclassified_all, 2)})で一致しました"
            f"(差={total_diff})。区分振替はカテゴリの訂正のみで、入金額の総額は変わりません。"
        )

    # 店舗全体売上・物販売上・施術/既存売上を明確に区別して報告する(2026-09-15確定。
    # product-audit-spec.md §17参照)。「物販売上」の金額を「店舗全体売上」と誤読しないため、
    # 月報集計シート自身が持つ実売合計(店舗全体)・内回数券等(施術・回数券)・内物販の
    # 内訳を読み取り、区分振替(§14)を反映した最終値も併せて出力する。
    monthly_summary = product_audit.parse_monthly_summary(wb)
    store_total_actual_sales = monthly_summary["store_total_actual_sales"]
    retail_sales_reported = monthly_summary["retail_sales_reported"]
    ticket_treatment_sales_reported = monthly_summary["ticket_treatment_sales_reported"]
    retail_sales_after_reclassification = round(total_engine, 2)
    ticket_treatment_sales_after_reclassification = (
        round(ticket_treatment_sales_reported + total_reclassified_all, 2)
        if ticket_treatment_sales_reported is not None else None
    )
    if store_total_actual_sales is not None:
        store_total_check_diff = round(
            retail_sales_after_reclassification
            + (ticket_treatment_sales_after_reclassification or 0)
            - store_total_actual_sales,
            2,
        )
        logger.info(
            f"店舗全体売上={store_total_actual_sales}円(月報集計シート実売実績累計・変更なし) / "
            f"物販売上(区分振替後)={retail_sales_after_reclassification}円 / "
            f"施術・既存売上(区分振替後)={ticket_treatment_sales_after_reclassification}円 "
            f"(差={store_total_check_diff})"
        )
        if abs(store_total_check_diff) >= 1.0:
            logger.error(
                "店舗全体売上と、物販売上+施術・既存売上の合計が一致しません。"
                "区分振替の反映漏れの可能性があるため要確認です。"
            )
    else:
        logger.error("月報集計シートから店舗全体の実売実績を読み取れませんでした(要確認)。")

    # --- 集計 ---
    classification_counts = {}
    for tx in all_transactions:
        classification_counts[tx.classification] = classification_counts.get(tx.classification, 0) + 1
    logger.info(f"判定件数: {classification_counts}")

    severity_counts = {"正常": 0, "既知差異": 0, "分類保留": 0, "注意": 0, "要確認": 0, "重大エラー": 0}
    for tx in all_transactions:
        severity_counts[tx.severity] = severity_counts.get(tx.severity, 0) + 1
    logger.info(f"重大度件数: {severity_counts}")

    # 価格監査(A軸)の状態別件数: 価格確定/価格要確認/既知差異。原価監査(B軸)とは独立に集計する
    # (2026-09-13確定。product-audit-spec.md §9参照)。
    price_status_counts = {"価格確定": 0, "価格要確認": 0, "既知差異": 0}
    for tx in all_transactions:
        price_status_counts[tx.price_status] = price_status_counts.get(tx.price_status, 0) + 1
    logger.info(f"価格監査 状態別件数: {price_status_counts}")

    def sum_attr(txs, attr):
        return round(sum(getattr(t, attr) or 0 for t in txs), 2)

    # 原価監査(B軸): 原価(仕入原価)自体が価格履歴に登録されているかどうかだけで判定する。
    # スタッフ社割価格が未登録(G判定)でも、原価が判明していれば原価確定として扱う
    # (2026-09-13確定。以前はprofit_confirmedとして価格の一致状況と一体で扱っていたため、
    # 在庫帳ベースの原価合計と食い違いが生じていた)。
    cost_confirmed_txs = [t for t in all_transactions if t.cost_confirmed]
    cost_unconfirmed_txs = [t for t in all_transactions if not t.cost_confirmed]

    # 監査結果の観点別件数(2026-09-15確定。product-audit-spec.md §20参照)。
    # 1取引が複数の観点に同時に該当しうるため、互いに排他的な分類ではなく、
    # それぞれ独立に集計する(例:価格異常かつ購入者不明、という取引もありうる)。
    price_anomaly_txs = [t for t in all_transactions if t.classification in ("C", "D")]
    unknown_product_txs = [t for t in all_transactions if t.product_status == "商品不明"]
    unknown_purchaser_txs = [t for t in all_transactions if t.purchaser_type == "購入者不明"]
    unresolved_share_txs = [
        t for t in all_transactions if t.transaction_type in ("shared_sale_candidate", "unknown")
    ]
    other_pending_txs = [
        t for t in all_transactions
        if t.severity == "要確認"
        and t.classification not in ("C", "D")
        and t.product_status != "商品不明"
        and t.transaction_type not in ("shared_sale_candidate", "unknown")
    ]
    audit_status_counts = {
        "price_anomaly": len(price_anomaly_txs),
        "unknown_product": len(unknown_product_txs),
        "unknown_purchaser": len(unknown_purchaser_txs),
        "unresolved_shared_sale": len(unresolved_share_txs),
        "cost_unconfirmed": len(cost_unconfirmed_txs),
        "other_pending_review": len(other_pending_txs),
    }
    logger.info(
        f"監査結果の観点別件数: 価格異常={audit_status_counts['price_anomaly']}件 / "
        f"商品不明={audit_status_counts['unknown_product']}件 / "
        f"購入者不明={audit_status_counts['unknown_purchaser']}件 / "
        f"売上シェア未解決={audit_status_counts['unresolved_shared_sale']}件 / "
        f"原価未確認={audit_status_counts['cost_unconfirmed']}件 / "
        f"その他要確認={audit_status_counts['other_pending_review']}件"
    )

    # 調査2: 「購入者区分」(スタッフ/社内購入/一般顧客/購入者不明)と「商品特定状態」
    # (商品特定済み/商品不明)を、それぞれ独立の軸として集計する(2026-09-14確定)。
    # 商品不明だからといって購入者区分の集計から除外せず、逆も行わない。
    # 1つの取引が両方の属性(例:購入者=一般顧客、商品=商品不明)を同時に持てる。
    # 各軸それぞれの合計が物販総売上と一致することを検証する。
    total_revenue = sum_attr(all_transactions, "tax_excl_revenue")

    purchaser_types = ["スタッフ", "社内購入", "一般顧客", "購入者不明"]
    purchaser_revenue = {
        p: sum_attr([t for t in all_transactions if t.purchaser_type == p], "tax_excl_revenue")
        for p in purchaser_types
    }
    purchaser_sum = round(sum(purchaser_revenue.values()), 2)
    purchaser_check_diff = round(purchaser_sum - total_revenue, 2)
    if abs(purchaser_check_diff) >= 1.0:
        logger.error(
            f"購入者区分別合計({purchaser_sum})が物販総売上({total_revenue})と一致しません"
            f"(差={purchaser_check_diff})。この監査は完成扱いにできません。"
        )
    else:
        logger.info(f"購入者区分別合計は物販総売上と一致しました(差={purchaser_check_diff})。")

    product_statuses = ["商品特定済み", "商品不明"]
    product_status_revenue = {
        s: sum_attr([t for t in all_transactions if t.product_status == s], "tax_excl_revenue")
        for s in product_statuses
    }
    product_status_sum = round(sum(product_status_revenue.values()), 2)
    product_status_check_diff = round(product_status_sum - total_revenue, 2)
    if abs(product_status_check_diff) >= 1.0:
        logger.error(
            f"商品特定状態別合計({product_status_sum})が物販総売上({total_revenue})と一致しません"
            f"(差={product_status_check_diff})。この監査は完成扱いにできません。"
        )
    else:
        logger.info(f"商品特定状態別合計は物販総売上と一致しました(差={product_status_check_diff})。")

    # 調査3: 在庫帳(店舗スタッフによる手入力の日別販売数量・原価集計)と、監査エンジンが
    # 日別シートの取引明細から積み上げた商品ごとの数量・原価を突き合わせる。
    # 「数量と在庫減少の不整合」「販売数量×原価と販売原価の不整合」の検出に使う。
    # 在庫帳のこれらの列は外部参照を経由しない(product-audit-spec.md §7参照)。
    INVENTORY_QTY_TOLERANCE = 0.01
    INVENTORY_COST_TOLERANCE_YEN = 1.0

    engine_qty_by_product: dict[str, float] = {}
    engine_cost_by_product: dict[str, float] = {}
    for t in all_transactions:
        if t.quantity:
            engine_qty_by_product[t.product_name] = engine_qty_by_product.get(t.product_name, 0.0) + t.quantity
        if t.cost_excl_tax_total is not None:
            engine_cost_by_product[t.product_name] = engine_cost_by_product.get(t.product_name, 0.0) + t.cost_excl_tax_total

    inventory_blocks = product_audit.parse_inventory_ledger(wb)
    inventory_reconciliation_rows = []
    for blk in inventory_blocks:
        name = blk["product_name"]
        if not name or name == "無":
            continue
        ledger_qty = blk["monthly_qty_sold"]
        ledger_cost_value = blk["monthly_cost_value"]
        engine_qty = engine_qty_by_product.get(name)
        engine_cost = engine_cost_by_product.get(name)
        if ledger_qty is None and engine_qty is None:
            continue
        qty_diff = (
            round((engine_qty or 0.0) - (ledger_qty or 0.0), 2)
            if not (ledger_qty is None and engine_qty is None) else None
        )
        cost_diff = (
            round((engine_cost or 0.0) - (ledger_cost_value or 0.0), 2)
            if not (ledger_cost_value is None and engine_cost is None) else None
        )
        qty_mismatch = qty_diff is not None and abs(qty_diff) > INVENTORY_QTY_TOLERANCE
        cost_mismatch = cost_diff is not None and abs(cost_diff) > INVENTORY_COST_TOLERANCE_YEN
        if not qty_mismatch and not cost_mismatch:
            continue
        severity = "要確認" if (qty_mismatch or cost_mismatch) else "正常"
        reasons = []
        if qty_mismatch:
            reasons.append("数量と在庫減少の不整合")
        if cost_mismatch:
            reasons.append("販売数量×原価と販売原価の不整合")
        inventory_reconciliation_rows.append({
            "product_name": name,
            "inventory_ledger_qty_sold": ledger_qty,
            "engine_qty_sold": engine_qty,
            "qty_diff": qty_diff,
            "inventory_ledger_cost_value": ledger_cost_value,
            "engine_cost_value": engine_cost,
            "cost_diff": cost_diff,
            "severity": severity,
            "reason": "・".join(reasons),
        })
    inventory_recon_path = reconciliation_out_dir / f"{store_id}_inventory_reconciliation.csv"
    io_utils.write_csv(
        inventory_recon_path, inventory_reconciliation_rows,
        fieldnames=["product_name", "inventory_ledger_qty_sold", "engine_qty_sold", "qty_diff",
                    "inventory_ledger_cost_value", "engine_cost_value", "cost_diff", "severity", "reason"],
    )
    logger.info(
        f"在庫帳との数量・原価突合を出力: {inventory_recon_path}"
        f"(不整合{len(inventory_reconciliation_rows)}件/商品ブロック{len(inventory_blocks)}件中)"
    )

    summary = {
        "store_id": store_id, "year_month": year_month,
        "source_file": str(source_path), "source_file_sha256": source_hash,
        "transaction_count": len(all_transactions),
        "classification_counts": classification_counts,
        "severity_counts": severity_counts,
        "price_status_counts": price_status_counts,
        "audit_status_counts": audit_status_counts,
        "inventory_reconciliation": {
            "product_blocks_checked": len(inventory_blocks),
            "mismatch_count": len(inventory_reconciliation_rows),
            "output_csv": str(inventory_recon_path),
        },
        "shared_sales": {
            "confirmed_group_count": len(confirmed_share_groups),
            "confirmed_2way_group_count": len(confirmed_2way_groups),
            "confirmed_3way_group_count": len(confirmed_3way_groups),
            "confirmed_row_count": sum(g["share_count"] for g in confirmed_share_groups),
            "candidate_group_count": len(candidate_share_groups),
            "candidate_row_count": sum(g["share_count"] for g in candidate_share_groups),
            "output_csv": str(shared_sales_path),
        },
        "category_reclassifications": {
            # 物販監査から除外し、施術/回数券(既存)側の売上へ振替した行(§14)。
            "count": len(all_reclassified_rows),
            "total_tax_excl_revenue": total_reclassified_all,
            "output_csv": str(category_reclass_path_out),
        },
        "af54_reconciliation": {
            "sum_af54": round(total_af54, 2), "engine_total": round(total_engine, 2),
            "reclassified_total": total_reclassified_all,
            "diff": total_diff, "matches": abs(total_diff) < 1.0,
        },
        # 店舗全体売上・物販売上・施術/既存売上を明確に区別するための内訳(§17)。
        # 「revenue_excl_tax.total」はあくまで物販売上であり、店舗全体売上ではないことに注意。
        "store_overall_revenue": {
            "store_total_actual_sales": store_total_actual_sales,
            "retail_sales_reported_before_reclassification": retail_sales_reported,
            "ticket_treatment_sales_reported_before_reclassification": ticket_treatment_sales_reported,
            "retail_sales_after_reclassification": retail_sales_after_reclassification,
            "ticket_treatment_sales_after_reclassification": ticket_treatment_sales_after_reclassification,
            "note": (
                "store_total_actual_sales(月報集計シート実売実績累計)は区分振替の影響を受けない"
                "店舗全体の実売合計。retail_sales_after_reclassification + "
                "ticket_treatment_sales_after_reclassification が store_total_actual_sales と"
                "一致することを確認する。"
            ),
        },
        "revenue_excl_tax": {
            # この値は「物販売上」であり、店舗全体売上ではない(store_overall_revenue参照)。
            "total": total_revenue,
            "by_purchaser_type": {**purchaser_revenue, "check_diff": purchaser_check_diff},
            "by_product_status": {**product_status_revenue, "check_diff": product_status_check_diff},
        },
        "cost_excl_tax": {
            # 実績原価。在庫帳当月販売金額・日別数量×原価・在庫増減式の3方式が一致することを
            # 確認済み(product-audit-spec.md §9)。管理会計上の売上原価として使用する。
            "actual_total": sum_attr(cost_confirmed_txs, "cost_excl_tax_total"),
        },
        "gross_profit": {
            # 管理会計上の物販粗利益: 物販売上総額 - 実績売上原価。価格監査・商品特定の
            # 未確定があっても、売上自体を管理会計の集計から除外しない(2026-09-14確定)。
            "management_accounting": round(total_revenue - sum_attr(cost_confirmed_txs, "cost_excl_tax_total"), 2),
            # 商品特定済み粗利益(旧称:確定粗利益): 商品不明(原価不明)を除いた、
            # 原価が判明している取引範囲のみの粗利益。監査上の確認範囲を示す別指標であり、
            # 管理会計上の物販粗利益とは異なる(2026-09-14確定。product-audit-spec.md §9)。
            "product_identified_confirmed": sum_attr(cost_confirmed_txs, "gross_profit"),
        },
        "gross_margin": {
            "management_accounting": (
                round((total_revenue - sum_attr(cost_confirmed_txs, "cost_excl_tax_total")) / total_revenue, 4)
                if total_revenue else None
            ),
            "product_identified_confirmed": (
                round(sum_attr(cost_confirmed_txs, "gross_profit") / sum_attr(cost_confirmed_txs, "tax_excl_revenue"), 4)
                if sum_attr(cost_confirmed_txs, "tax_excl_revenue") else None
            ),
        },
        "cost_confirmed_transaction_count": len(cost_confirmed_txs),
        "cost_unconfirmed_transaction_count": len(cost_unconfirmed_txs),
        "revenue_excl_tax_cost_confirmed": sum_attr(cost_confirmed_txs, "tax_excl_revenue"),
        "revenue_excl_tax_cost_unconfirmed": sum_attr(cost_unconfirmed_txs, "tax_excl_revenue"),
    }

    out_dir = io_utils.OUTPUT_DIR / "product_audit" / year_month
    summary_path = out_dir / f"{store_id}_product_audit_summary.json"
    io_utils.write_json(summary_path, summary)
    logger.info(f"監査サマリを出力: {summary_path}")

    detail_rows = []
    for t in all_transactions:
        detail_rows.append({
            "date": t.date, "row": t.sheet_row,
            "customer_name": t.customer_name or "(空欄)",
            "purchaser_type": t.purchaser_type, "product_status": t.product_status,
            "product_name": t.product_name,
            "quantity": t.quantity,
            "regular_price_incl_tax_expected": t.regular_price_incl_tax_expected,
            "staff_price_incl_tax_expected": t.staff_price_incl_tax_expected,
            "actual_price_incl_tax": t.actual_price_incl_tax,
            "discount_incl_tax": t.discount_incl_tax,
            "tax_excl_revenue": t.tax_excl_revenue,
            "cost_excl_tax_total": t.cost_excl_tax_total,
            "gross_profit": t.gross_profit,
            "gross_margin": t.gross_margin,
            "classification": t.classification,
            "classification_detail": t.classification_detail,
            "price_status": t.price_status,
            "cost_confirmed": t.cost_confirmed,
            "severity": t.severity,
            "flags": "・".join(t.flags) if t.flags else "",
            "purchaser_alias_note": t.purchaser_alias_note,
            "purchaser_name_note": t.purchaser_name_note,
            "quantity_correction_note": t.quantity_correction_note,
            "transaction_type": t.transaction_type,
            "share_group_id": t.share_group_id or "",
            "share_count": t.share_count,
            "share_total": t.share_total,
            "linked_product": t.linked_product or "",
            "linked_rows": "・".join(str(r) for r in t.linked_rows) if t.linked_rows else "",
            "share_marker_raw": t.share_marker_raw or "",
        })
    detail_path = out_dir / f"{store_id}_product_audit_detail.csv"
    io_utils.write_csv(
        detail_path, detail_rows,
        fieldnames=["date", "row", "customer_name", "purchaser_type", "product_status", "product_name",
                    "quantity", "regular_price_incl_tax_expected", "staff_price_incl_tax_expected",
                    "actual_price_incl_tax", "discount_incl_tax", "tax_excl_revenue",
                    "cost_excl_tax_total", "gross_profit", "gross_margin",
                    "classification", "classification_detail", "price_status", "cost_confirmed",
                    "severity", "flags", "purchaser_alias_note", "purchaser_name_note",
                    "quantity_correction_note",
                    "transaction_type", "share_group_id", "share_count", "share_total",
                    "linked_product", "linked_rows", "share_marker_raw"],
    )
    logger.info(f"取引明細を出力: {detail_path}")

    # 既知差異一覧: 原因特定済みの表示価格変動(K判定)。要確認一覧には含めない
    # (2026-09-13確定。product-audit-spec.md §8参照)。
    known_discrepancy_rows = [r for r in detail_rows if r["severity"] == "既知差異"]
    known_discrepancy_path = out_dir / f"{store_id}_product_audit_known_discrepancies.csv"
    io_utils.write_csv(known_discrepancy_path, known_discrepancy_rows, fieldnames=detail_rows[0].keys() if detail_rows else [])
    logger.info(f"既知差異一覧を出力: {known_discrepancy_path}({len(known_discrepancy_rows)}件)")

    # 分類保留一覧: 事実関係は確定済みだが、社割ルール等のマスター登録が未了のもの(§15)。
    # 「既知差異」と同様、原因不明の要確認とは区別し、要確認一覧からは除外する。
    pending_classification_rows = [r for r in detail_rows if r["severity"] == "分類保留"]
    pending_classification_path = out_dir / f"{store_id}_product_audit_pending_classification.csv"
    io_utils.write_csv(
        pending_classification_path, pending_classification_rows,
        fieldnames=detail_rows[0].keys() if detail_rows else [],
    )
    logger.info(f"分類保留一覧を出力: {pending_classification_path}({len(pending_classification_rows)}件)")

    # 要確認一覧(P項目): 重大度が「注意/要確認/重大エラー」の取引を含める。
    # 「既知差異」「分類保留」は原因・事実関係が確定済みのため、要確認一覧からは明示的に除外する。
    # 判定区分がA(正常)であっても、追加フラグ(異常値引き・社割価格流用の可能性等)により
    # 重大度が引き上げられている取引を取りこぼさないため、classificationではなくseverityで判定する。
    error_rows = [r for r in detail_rows if r["severity"] not in ("正常", "既知差異", "分類保留")]
    error_path = out_dir / f"{store_id}_product_audit_errors.csv"
    io_utils.write_csv(error_path, error_rows, fieldnames=detail_rows[0].keys() if detail_rows else [])
    logger.info(f"要確認一覧(正常・既知差異・分類保留を除く)を出力: {error_path}({len(error_rows)}件)")

    unconfirmed_rows = [r for r in detail_rows if not r["cost_confirmed"]]
    unconfirmed_path = out_dir / f"{store_id}_product_audit_cost_unconfirmed.csv"
    io_utils.write_csv(unconfirmed_path, unconfirmed_rows, fieldnames=detail_rows[0].keys() if detail_rows else [])
    logger.info(f"原価未確定取引一覧を出力: {unconfirmed_path}({len(unconfirmed_rows)}件)")

    for cls in sorted(classification_counts):
        if cls == "K":
            logger.info(f"判定{cls}(既知差異): {classification_counts[cls]}件")
        elif cls == "S":
            logger.info(f"判定{cls}(売上シェア確定): {classification_counts[cls]}件")
        elif cls == "N":
            logger.info(f"判定{cls}(社内例外取引): {classification_counts[cls]}件")
        elif cls != "A":
            logger.warning(f"判定{cls}: {classification_counts[cls]}件")

    ledger_info = io_utils.update_ledger(
        store_id=f"{store_id}_product_audit", year_month=year_month,
        source_file_sha256=source_hash, source_file_path=str(source_path),
        output_paths=[str(summary_path), str(detail_path), str(error_path), str(unconfirmed_path),
                      str(known_discrepancy_path), str(pending_classification_path),
                      str(day_reconciliation_path), str(inventory_recon_path),
                      str(shared_sales_path), str(category_reclass_path_out)],
    )
    if ledger_info["is_rerun_same_source"]:
        logger.info("同一ソースファイルでの再実行です。出力は上書きされ、二重計上は発生していません。")

    logger.info("=== 完了 ===")
    logger.info(f"ログファイル: {log_path}")


if __name__ == "__main__":
    main()
